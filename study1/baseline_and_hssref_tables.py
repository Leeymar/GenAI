# -*- coding: utf-8 -*-
"""
生成两类补充回归表：
1) 基准回归表（不含调节变量）：5个知识扩散指标分别回归于 HSS_Continuous + controls + FE
2) HSS参考文献占比回归表：hss_reference_share ~ HSS_Continuous + controls + FE

默认读取：
E:\PythonProject\GenAI\data_process\derived\genai_regression_enriched_v3_fixed.csv

默认输出：
E:\PythonProject\GenAI\data_process\derived\baseline_and_hssref_tables.xlsx

说明：
- 与你现有表保持一致，使用 OLS + HC1 稳健标准误
- Hit_Rate_10_year 仍按线性概率模型（OLS）处理，以保持和现有回归表口径一致
- 生成一个 Excel，其中包含 pretty/raw/long/full_params/model_summary/readme 等 sheet
"""

from __future__ import annotations

import math
import os
from dataclasses import dataclass
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_CSV = r"E:\PythonProject\GenAI\data_process\derived\genai_regression_enriched_v3_fixed.csv"
OUTPUT_XLSX = r"E:\PythonProject\GenAI\data_process\derived\baseline_and_hssref_tables.xlsx"


DEP_VARS = [
    ("平均引用影响 (Log_RCR_p)", "Log_RCR", "知识扩散强度"),
    ("高影响论文概率 (Hit{10}_p)", "Hit_Rate_10_year", "知识扩散强度"),
    ("施引学科广度 (Citation_Field_Diversity_p)", "Citation_Field_Diversity", "知识扩散广度"),
    ("HSS施引占比 (HSS_Citing_Share_p)", "HSS_Citing_Share", "知识扩散广度"),
    ("施引学科熵 (Citation_Entropy_p)", "Citation_Entropy", "知识扩散均匀度"),
]

BASELINE_ROWS = [
    ("团队HSS浓度", "HSS_Continuous"),
    ("团队规模（对数）", "log_team_size"),
    ("先验知识存量（对数）", "log_prior_knowledge"),
    ("常数项", "Intercept"),
]

HSSREF_ROW_ORDER = BASELINE_ROWS

CONTROLS = ["log_team_size", "log_prior_knowledge"]
KEY_X = "HSS_Continuous"
YEAR_FE = "year"
TOPIC_FE = "arxiv_primary_category"
ROBUST_COV = "HC1"


@dataclass
class ModelResult:
    model_group: str
    dep_label: str
    dep_col: str
    spec_label: str
    formula: str
    nobs: int
    r2: float
    adj_r2: float
    params_df: pd.DataFrame


def stars(p: float) -> str:
    if pd.isna(p):
        return ""
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.1:
        return "*"
    return ""


def fmt_coef_se(coef: float, se: float, p: float, digits: int = 4) -> str:
    if pd.isna(coef):
        return ""
    return f"{coef:.{digits}f}{stars(p)}\n({se:.{digits}f})"


def safe_float(x):
    try:
        return float(x)
    except Exception:
        return np.nan


def fit_ols(df: pd.DataFrame, dep: str, add_topic_fe: bool) -> Tuple[ModelResult, pd.DataFrame]:
    needed = [dep, KEY_X, *CONTROLS, YEAR_FE]
    if add_topic_fe:
        needed.append(TOPIC_FE)
    sub = df[needed].copy().dropna()

    fe_terms = [f"C({YEAR_FE})"]
    if add_topic_fe:
        fe_terms.append(f"C({TOPIC_FE})")

    formula = f"{dep} ~ {KEY_X} + {' + '.join(CONTROLS)} + " + " + ".join(fe_terms)
    res = smf.ols(formula, data=sub).fit(cov_type=ROBUST_COV)

    rows = []
    for pname in res.params.index:
        rows.append({
            "参数名": pname,
            "系数": safe_float(res.params.get(pname)),
            "标准误": safe_float(res.bse.get(pname)),
            "p值": safe_float(res.pvalues.get(pname)),
            "显著性": stars(safe_float(res.pvalues.get(pname))),
        })
    params_df = pd.DataFrame(rows)

    model = ModelResult(
        model_group="基准回归" if dep != "hss_reference_share" else "HSS参考文献占比回归",
        dep_label=dep,
        dep_col=dep,
        spec_label="有主题FE" if add_topic_fe else "无主题FE",
        formula=formula,
        nobs=int(res.nobs),
        r2=safe_float(res.rsquared),
        adj_r2=safe_float(res.rsquared_adj),
        params_df=params_df,
    )
    return model, sub


def extract_core_rows(model: ModelResult, dep_label: str, dep_col: str, display_label: str) -> pd.DataFrame:
    df = model.params_df.copy()
    rename_map = {
        "HSS_Continuous": "团队HSS浓度",
        "log_team_size": "团队规模（对数）",
        "log_prior_knowledge": "先验知识存量（对数）",
        "Intercept": "常数项",
    }
    df["变量"] = df["参数名"].map(rename_map)
    df = df[df["变量"].notna()].copy()
    df["格式化"] = df.apply(lambda r: fmt_coef_se(r["系数"], r["标准误"], r["p值"]), axis=1)
    df["被解释变量"] = display_label
    df["被解释变量列名"] = dep_col
    df["规格"] = model.spec_label
    df["观测值"] = model.nobs
    df["R2"] = model.r2
    df["调整后R2"] = model.adj_r2
    df["年份固定效应"] = "是"
    df["arXiv主分类固定效应"] = "是" if model.spec_label == "有主题FE" else "否"
    df["公式"] = model.formula
    keep_cols = [
        "被解释变量", "被解释变量列名", "规格", "变量", "参数名", "系数", "标准误", "p值", "显著性",
        "格式化", "观测值", "R2", "调整后R2", "年份固定效应", "arXiv主分类固定效应", "公式"
    ]
    return df[keep_cols]


def build_baseline_pretty(core_long: pd.DataFrame) -> pd.DataFrame:
    rows = [r[0] for r in BASELINE_ROWS] + ["观测值", "R²", "调整后R²", "年份固定效应", "arXiv主分类固定效应"]
    cols = []
    for disp, dep_col, group in DEP_VARS:
        cols.append((group, disp, "无主题FE"))
        cols.append((group, disp, "有主题FE"))

    matrix = []
    for row_name in rows:
        row = {"变量": row_name}
        for _, disp, spec in cols:
            sub = core_long[(core_long["被解释变量"] == disp) & (core_long["规格"] == spec)]
            if row_name in [r[0] for r in BASELINE_ROWS]:
                val = sub.loc[sub["变量"] == row_name, "格式化"]
                row[(disp, spec)] = val.iloc[0] if len(val) else ""
            elif row_name == "观测值":
                row[(disp, spec)] = str(int(sub["观测值"].iloc[0])) if len(sub) else ""
            elif row_name == "R²":
                row[(disp, spec)] = f"{sub['R2'].iloc[0]:.4f}" if len(sub) else ""
            elif row_name == "调整后R²":
                row[(disp, spec)] = f"{sub['调整后R2'].iloc[0]:.4f}" if len(sub) else ""
            elif row_name == "年份固定效应":
                row[(disp, spec)] = sub["年份固定效应"].iloc[0] if len(sub) else ""
            elif row_name == "arXiv主分类固定效应":
                row[(disp, spec)] = sub["arXiv主分类固定效应"].iloc[0] if len(sub) else ""
        matrix.append(row)

    out = pd.DataFrame(matrix)
    ordered_cols = ["变量"]
    for _, disp, spec in cols:
        ordered_cols.append((disp, spec))
    out = out[ordered_cols]

    header0 = ["变量"]
    header1 = [""]
    header2 = [""]
    for group, disp, spec in cols:
        header0.append(group)
        header1.append(disp)
        header2.append(spec)

    data_rows = []
    for _, r in out.iterrows():
        data_rows.append([r["变量"]] + [r[(disp, spec)] for _, disp, spec in cols])
    pretty = pd.DataFrame([header0, header1, header2] + data_rows)
    return pretty


def build_hssref_pretty(core_long: pd.DataFrame) -> pd.DataFrame:
    # 这里只处理因变量 hss_reference_share
    rows = [r[0] for r in HSSREF_ROW_ORDER] + ["观测值", "R²", "调整后R²", "年份固定效应", "arXiv主分类固定效应"]
    result = []
    for rn in rows:
        row = {"变量": rn}
        for spec in ["无主题FE", "有主题FE"]:
            sub = core_long[core_long["规格"] == spec]
            if rn in [r[0] for r in HSSREF_ROW_ORDER]:
                val = sub.loc[sub["变量"] == rn, "格式化"]
                row[spec] = val.iloc[0] if len(val) else ""
            elif rn == "观测值":
                row[spec] = str(int(sub["观测值"].iloc[0])) if len(sub) else ""
            elif rn == "R²":
                row[spec] = f"{sub['R2'].iloc[0]:.4f}" if len(sub) else ""
            elif rn == "调整后R²":
                row[spec] = f"{sub['调整后R2'].iloc[0]:.4f}" if len(sub) else ""
            elif rn == "年份固定效应":
                row[spec] = sub["年份固定效应"].iloc[0] if len(sub) else ""
            elif rn == "arXiv主分类固定效应":
                row[spec] = sub["arXiv主分类固定效应"].iloc[0] if len(sub) else ""
        result.append(row)

    df = pd.DataFrame(result)
    pretty = pd.DataFrame([
        ["变量", "HSS参考文献占比 (hss_reference_share)", ""],
        ["", "无主题FE", "有主题FE"],
    ] + df[["变量", "无主题FE", "有主题FE"]].values.tolist())
    return pretty


def make_full_params(model: ModelResult, display_label: str, dep_col: str) -> pd.DataFrame:
    df = model.params_df.copy()
    df.insert(0, "模型组", model.model_group)
    df.insert(1, "被解释变量", display_label)
    df.insert(2, "被解释变量列名", dep_col)
    df.insert(3, "规格", model.spec_label)
    df["观测值"] = model.nobs
    df["R2"] = model.r2
    df["调整后R2"] = model.adj_r2
    df["年份固定效应"] = "是"
    df["arXiv主分类固定效应"] = "是" if model.spec_label == "有主题FE" else "否"
    df["公式"] = model.formula
    return df


def auto_format_xlsx(path: str) -> None:
    wb = load_workbook(path)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    panel_fill = PatternFill("solid", fgColor="EAF2F8")

    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
                cell.border = border

        max_row = ws.max_row
        max_col = ws.max_column
        for r in range(1, min(max_row, 3) + 1):
            for c in range(1, max_col + 1):
                ws.cell(r, c).font = Font(bold=True)
                ws.cell(r, c).fill = header_fill

        # 调整列宽
        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            max_len = 0
            for cell in col_cells[:120]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(val), 40))
            ws.column_dimensions[letter].width = max(10, min(max_len + 2, 28))

        # 第一列宽一点
        ws.column_dimensions["A"].width = 24

        # 针对 pretty sheet 加强表头
        if ws.title in {"baseline_pretty", "hssref_pretty"}:
            for c in range(1, ws.max_column + 1):
                ws.cell(1, c).fill = panel_fill
                ws.cell(1, c).font = Font(bold=True)
                ws.cell(2, c).font = Font(bold=True)
            ws.freeze_panes = "B4" if ws.title == "baseline_pretty" else "B3"

    wb.save(path)


def main() -> None:
    input_csv = os.environ.get("GENAI_INPUT_CSV", INPUT_CSV)
    output_xlsx = os.environ.get("GENAI_OUTPUT_XLSX", OUTPUT_XLSX)

    if not os.path.exists(input_csv):
        raise FileNotFoundError(
            f"未找到输入 CSV：{input_csv}\n"
            f"请把环境变量 GENAI_INPUT_CSV 指向正确文件，或修改脚本顶部 INPUT_CSV。"
        )

    df = pd.read_csv(input_csv, low_memory=False)

    baseline_core_list: List[pd.DataFrame] = []
    baseline_full_list: List[pd.DataFrame] = []
    baseline_summary_rows: List[Dict] = []

    for disp, dep_col, group in DEP_VARS:
        for add_topic in [False, True]:
            model, _ = fit_ols(df, dep_col, add_topic)
            baseline_core_list.append(extract_core_rows(model, dep_col, dep_col, disp))
            baseline_full_list.append(make_full_params(model, disp, dep_col))
            baseline_summary_rows.append({
                "模型组": "基准回归",
                "被解释变量": disp,
                "被解释变量列名": dep_col,
                "规格": model.spec_label,
                "观测值": model.nobs,
                "R2": model.r2,
                "调整后R2": model.adj_r2,
                "年份固定效应": "是",
                "arXiv主分类固定效应": "是" if add_topic else "否",
                "公式": model.formula,
            })

    baseline_core_long = pd.concat(baseline_core_list, ignore_index=True)
    baseline_full_params = pd.concat(baseline_full_list, ignore_index=True)
    baseline_model_summary = pd.DataFrame(baseline_summary_rows)
    baseline_pretty = build_baseline_pretty(baseline_core_long)

    # 原始宽表
    raw_rows = [r[0] for r in BASELINE_ROWS] + ["观测值", "R²", "调整后R²", "年份固定效应", "arXiv主分类固定效应"]
    baseline_raw = pd.DataFrame({"变量": raw_rows})
    for disp, dep_col, group in DEP_VARS:
        for spec in ["无主题FE", "有主题FE"]:
            sub = baseline_core_long[(baseline_core_long["被解释变量"] == disp) & (baseline_core_long["规格"] == spec)]
            vals = []
            for rn in raw_rows:
                if rn in [r[0] for r in BASELINE_ROWS]:
                    x = sub.loc[sub["变量"] == rn, "格式化"]
                    vals.append(x.iloc[0] if len(x) else "")
                elif rn == "观测值":
                    vals.append(str(int(sub["观测值"].iloc[0])) if len(sub) else "")
                elif rn == "R²":
                    vals.append(f"{sub['R2'].iloc[0]:.4f}" if len(sub) else "")
                elif rn == "调整后R²":
                    vals.append(f"{sub['调整后R2'].iloc[0]:.4f}" if len(sub) else "")
                elif rn == "年份固定效应":
                    vals.append(sub["年份固定效应"].iloc[0] if len(sub) else "")
                elif rn == "arXiv主分类固定效应":
                    vals.append(sub["arXiv主分类固定效应"].iloc[0] if len(sub) else "")
            baseline_raw[f"{disp}\n{spec}"] = vals

    # hss_reference_share 作为因变量
    hssref_core_list: List[pd.DataFrame] = []
    hssref_full_list: List[pd.DataFrame] = []
    hssref_summary_rows: List[Dict] = []
    hssref_display = "HSS参考文献占比 (hss_reference_share)"
    hssref_dep = "hss_reference_share"

    for add_topic in [False, True]:
        model, _ = fit_ols(df, hssref_dep, add_topic)
        hssref_core_list.append(extract_core_rows(model, hssref_dep, hssref_dep, hssref_display))
        hssref_full_list.append(make_full_params(model, hssref_display, hssref_dep))
        hssref_summary_rows.append({
            "模型组": "HSS参考文献占比回归",
            "被解释变量": hssref_display,
            "被解释变量列名": hssref_dep,
            "规格": model.spec_label,
            "观测值": model.nobs,
            "R2": model.r2,
            "调整后R2": model.adj_r2,
            "年份固定效应": "是",
            "arXiv主分类固定效应": "是" if add_topic else "否",
            "公式": model.formula,
        })

    hssref_core_long = pd.concat(hssref_core_list, ignore_index=True)
    hssref_full_params = pd.concat(hssref_full_list, ignore_index=True)
    hssref_model_summary = pd.DataFrame(hssref_summary_rows)
    hssref_pretty = build_hssref_pretty(hssref_core_long)

    readme = pd.DataFrame([
        ["输入文件", input_csv],
        ["输出文件", output_xlsx],
        ["模型1", "基准回归：Y ~ HSS_Continuous + log_team_size + log_prior_knowledge + C(year) [+ C(arxiv_primary_category)]"],
        ["模型2", "HSS参考文献占比回归：hss_reference_share ~ HSS_Continuous + log_team_size + log_prior_knowledge + C(year) [+ C(arxiv_primary_category)]"],
        ["团队HSS浓度变量", "HSS_Continuous"],
        ["控制变量1", "log_team_size"],
        ["控制变量2", "log_prior_knowledge"],
        ["年份固定效应变量", YEAR_FE],
        ["arXiv主分类固定效应变量", TOPIC_FE],
        ["稳健标准误", ROBUST_COV],
        ["说明", "Hit_Rate_10_year 仍按 OLS 线性概率模型处理，以保持与现有回归表一致"],
        ["说明", "pretty sheet 为可直接粘贴论文的回归系数表；core_long/full_params/model_summary 为核对明细"],
        [np.nan, np.nan],
        ["因变量显示名", "CSV 实际列名"],
        ["平均引用影响 (Log_RCR_p)", "Log_RCR"],
        ["高影响论文概率 (Hit{10}_p)", "Hit_Rate_10_year"],
        ["施引学科广度 (Citation_Field_Diversity_p)", "Citation_Field_Diversity"],
        ["HSS施引占比 (HSS_Citing_Share_p)", "HSS_Citing_Share"],
        ["施引学科熵 (Citation_Entropy_p)", "Citation_Entropy"],
        ["HSS参考文献占比 (hss_reference_share)", "hss_reference_share"],
    ], columns=["项目", "内容"])

    os.makedirs(os.path.dirname(output_xlsx), exist_ok=True)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="readme", index=False)
        baseline_pretty.to_excel(writer, sheet_name="baseline_pretty", index=False, header=False)
        baseline_raw.to_excel(writer, sheet_name="baseline_raw", index=False)
        baseline_core_long.to_excel(writer, sheet_name="baseline_core_long", index=False)
        baseline_full_params.to_excel(writer, sheet_name="baseline_full_params", index=False)
        baseline_model_summary.to_excel(writer, sheet_name="baseline_model_summary", index=False)
        hssref_pretty.to_excel(writer, sheet_name="hssref_pretty", index=False, header=False)
        hssref_core_long.to_excel(writer, sheet_name="hssref_core_long", index=False)
        hssref_full_params.to_excel(writer, sheet_name="hssref_full_params", index=False)
        hssref_model_summary.to_excel(writer, sheet_name="hssref_model_summary", index=False)

    auto_format_xlsx(output_xlsx)
    print(f"已输出：{output_xlsx}")


if __name__ == "__main__":
    main()
