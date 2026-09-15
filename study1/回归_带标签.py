
# -*- coding: utf-8 -*-
"""
3.5 双规格回归（修正版）
------------------------------------------------------------
按用户最终确认口径：
- 因变量 Y：5 个“知识扩散”指标
    1) Log_RCR_p                     -> 实际列名: Log_RCR
    2) Hit{10}_p                     -> 实际列名: Hit_Rate_10_year
    3) Citation_Field_Diversity_p    -> 实际列名: Citation_Field_Diversity
    4) HSS_Citing_Share_p            -> 实际列名: HSS_Citing_Share
    5) Citation_Entropy_p            -> 实际列名: Citation_Entropy
- 调节变量 Integration_p：4 个“参考文献侧”指标
    1) reference_field_diversity
    2) reference_entropy
    3) hss_reference_share
    4) cs_core_reference_share
- 控制变量 X_p：
    1) log_team_size
    2) log_prior_knowledge
- 固定效应：
    1) 年份固定效应 C(year)
    2) arXiv 主分类固定效应 C(arxiv_primary_category)

模型：
Y = α + β1*HSS_Continuous + β2*Integration + β3*(HSS_Continuous*Integration)
    + γ1*log_team_size + γ2*log_prior_knowledge + year FE [+ topic FE] + ε

输出：
1) readme
2) big_table_pretty   -> 论文用大表（3 行抬头，双规格对照）
3) big_table_raw      -> 原始宽表
4) core_long          -> 核心系数长表
5) full_params        -> 完整参数表（含固定效应虚拟变量）
6) model_summary      -> 模型摘要
"""

from pathlib import Path
import re
import warnings

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# =========================================================
# 1. 路径设置
# =========================================================
CSV_PATH = Path(r"E:\PythonProject\GenAI\data_process\derived\genai_regression_enriched_v3_fixed.csv")
OUTPUT_XLSX = CSV_PATH.with_name("table_3_5_topic_fe_compare_FIXED_5Y.xlsx")

ROBUST_SE = "HC1"


# =========================================================
# 2. 精确变量映射（按你最后确认的口径）
# =========================================================

# 核心解释变量
BASE_VAR = "HSS_Continuous"

# 控制变量
CONTROL_VARS = [
    "log_team_size",
    "log_prior_knowledge",
]

# 固定效应
YEAR_FE_VAR = "year"
TOPIC_FE_VAR = "arxiv_primary_category"

# 因变量：显示名 -> 实际列名
OUTCOME_MAP = {
    "平均引用影响 (Log_RCR_p)": "Log_RCR",
    "高影响论文概率 (Hit{10}_p)": "Hit_Rate_10_year",
    "施引学科广度 (Citation_Field_Diversity_p)": "Citation_Field_Diversity",
    "HSS施引占比 (HSS_Citing_Share_p)": "HSS_Citing_Share",
    "施引学科熵 (Citation_Entropy_p)": "Citation_Entropy",
}

# 因变量所属维度（用于大表抬头）
OUTCOME_GROUP_MAP = {
    "平均引用影响 (Log_RCR_p)": "知识扩散强度",
    "高影响论文概率 (Hit{10}_p)": "知识扩散强度",
    "施引学科广度 (Citation_Field_Diversity_p)": "知识扩散广度",
    "HSS施引占比 (HSS_Citing_Share_p)": "知识扩散广度",
    "施引学科熵 (Citation_Entropy_p)": "知识扩散均匀度",
}

# 调节变量：显示名 -> 实际列名
MODERATOR_MAP = {
    "参考文献学科数量": "reference_field_diversity",
    "参考文献学科熵": "reference_entropy",
    "HSS参考文献占比": "hss_reference_share",
    "CS核心参考文献占比": "cs_core_reference_share",
}

SPEC_LIST = [
    ("无主题FE", False),
    ("有主题FE", True),
]


# =========================================================
# 3. 工具函数
# =========================================================

def safe_name(col: str) -> str:
    """公式中安全引用变量名。"""
    if re.match(r"^[A-Za-z_]\w*$", col):
        return col
    return f'Q("{col}")'


def star(p):
    if pd.isna(p):
        return ""
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.10:
        return "*"
    return ""


def fmt_coef_cell(model, term):
    coef = model.params.get(term, np.nan)
    se = model.bse.get(term, np.nan)
    p = model.pvalues.get(term, np.nan)
    if pd.isna(coef):
        return ""
    return f"{coef:.4f}{star(p)}\n({se:.4f})"


def build_formula(y_col, moderator_col, add_topic_fe):
    rhs = [
        f"{safe_name(BASE_VAR)} * {safe_name(moderator_col)}",
        safe_name(CONTROL_VARS[0]),
        safe_name(CONTROL_VARS[1]),
        f"C({safe_name(YEAR_FE_VAR)})",
    ]
    if add_topic_fe:
        rhs.append(f"C({safe_name(TOPIC_FE_VAR)})")
    return f"{safe_name(y_col)} ~ " + " + ".join(rhs)


def find_interaction_term(base_var, moderator_col, params_index):
    t1 = f"{safe_name(base_var)}:{safe_name(moderator_col)}"
    t2 = f"{safe_name(moderator_col)}:{safe_name(base_var)}"
    if t1 in params_index:
        return t1
    if t2 in params_index:
        return t2
    return t1


def ensure_required_columns(df, required_cols):
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        raise KeyError(
            "CSV 缺少以下必要列：\n- " + "\n- ".join(missing)
        )


def set_common_style(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)


# =========================================================
# 4. 读取数据
# =========================================================
if not CSV_PATH.exists():
    raise FileNotFoundError(f"找不到文件：{CSV_PATH}")

try:
    df = pd.read_csv(CSV_PATH, low_memory=False, encoding="utf-8-sig")
except UnicodeDecodeError:
    df = pd.read_csv(CSV_PATH, low_memory=False)

print(f"已读取数据：{CSV_PATH}")
print(f"数据维度：{df.shape[0]} 行 × {df.shape[1]} 列")


# =========================================================
# 5. 检查列是否齐全
# =========================================================
required_cols = [
    BASE_VAR,
    YEAR_FE_VAR,
    TOPIC_FE_VAR,
    *CONTROL_VARS,
    *OUTCOME_MAP.values(),
    *MODERATOR_MAP.values(),
]
ensure_required_columns(df, required_cols)

# 数值型转换
numeric_cols = [
    BASE_VAR,
    *CONTROL_VARS,
    *OUTCOME_MAP.values(),
    *MODERATOR_MAP.values(),
    YEAR_FE_VAR,
]
for c in numeric_cols:
    df[c] = pd.to_numeric(df[c], errors="coerce")

# 主题变量转字符串
df[TOPIC_FE_VAR] = df[TOPIC_FE_VAR].astype("string")


# =========================================================
# 6. 回归
# =========================================================
model_store = {}
core_rows = []
full_rows = []
summary_rows = []

for moderator_label, moderator_col in MODERATOR_MAP.items():
    for outcome_label, outcome_col in OUTCOME_MAP.items():

        # 关键：两套规格共用同一批样本，保证可比
        sample_cols = [
            outcome_col,
            BASE_VAR,
            moderator_col,
            *CONTROL_VARS,
            YEAR_FE_VAR,
            TOPIC_FE_VAR,
        ]
        dsub = df[sample_cols].dropna().copy()

        if dsub.empty:
            print(f"[跳过] {outcome_label} × {moderator_label}：样本为空")
            continue

        for spec_name, add_topic_fe in SPEC_LIST:
            formula = build_formula(outcome_col, moderator_col, add_topic_fe)
            model = smf.ols(formula=formula, data=dsub).fit(cov_type=ROBUST_SE)

            interaction_term = find_interaction_term(BASE_VAR, moderator_col, model.params.index)

            terms_map = {
                "团队HSS浓度": safe_name(BASE_VAR),
                moderator_label: safe_name(moderator_col),
                f"团队HSS浓度 × {moderator_label}": interaction_term,
                "团队规模（对数）": safe_name(CONTROL_VARS[0]),
                "先验知识存量（对数）": safe_name(CONTROL_VARS[1]),
                "常数项": "Intercept",
            }

            model_store[(moderator_label, outcome_label, spec_name)] = {
                "model": model,
                "nobs": int(model.nobs),
                "r2": model.rsquared,
                "adj_r2": model.rsquared_adj,
                "year_fe": "是",
                "topic_fe": "是" if add_topic_fe else "否",
                "formula": formula,
                "terms_map": terms_map,
                "outcome_col": outcome_col,
                "moderator_col": moderator_col,
            }

            # 核心系数长表
            for row_label, term in terms_map.items():
                coef = model.params.get(term, np.nan)
                se = model.bse.get(term, np.nan)
                pval = model.pvalues.get(term, np.nan)

                core_rows.append({
                    "调节变量": moderator_label,
                    "调节变量列名": moderator_col,
                    "被解释变量": outcome_label,
                    "被解释变量列名": outcome_col,
                    "规格": spec_name,
                    "变量": row_label,
                    "参数名": term,
                    "系数": coef,
                    "标准误": se,
                    "p值": pval,
                    "显著性": star(pval),
                    "格式化": f"{coef:.4f}{star(pval)} ({se:.4f})" if pd.notna(coef) else "",
                    "观测值": int(model.nobs),
                    "R2": model.rsquared,
                    "调整后R2": model.rsquared_adj,
                    "年份固定效应": "是",
                    "arXiv主分类固定效应": "是" if add_topic_fe else "否",
                    "公式": formula,
                })

            # 完整参数表
            for term in model.params.index:
                full_rows.append({
                    "调节变量": moderator_label,
                    "调节变量列名": moderator_col,
                    "被解释变量": outcome_label,
                    "被解释变量列名": outcome_col,
                    "规格": spec_name,
                    "参数名": term,
                    "系数": model.params.get(term, np.nan),
                    "标准误": model.bse.get(term, np.nan),
                    "p值": model.pvalues.get(term, np.nan),
                    "显著性": star(model.pvalues.get(term, np.nan)),
                    "观测值": int(model.nobs),
                    "R2": model.rsquared,
                    "调整后R2": model.rsquared_adj,
                    "年份固定效应": "是",
                    "arXiv主分类固定效应": "是" if add_topic_fe else "否",
                    "公式": formula,
                })

            # 模型摘要
            summary_rows.append({
                "调节变量": moderator_label,
                "调节变量列名": moderator_col,
                "被解释变量": outcome_label,
                "被解释变量列名": outcome_col,
                "规格": spec_name,
                "观测值": int(model.nobs),
                "R2": model.rsquared,
                "调整后R2": model.rsquared_adj,
                "年份固定效应": "是",
                "arXiv主分类固定效应": "是" if add_topic_fe else "否",
                "公式": formula,
            })

        print(f"[完成] {outcome_label} × {moderator_label}")


# =========================================================
# 7. 生成原始大表（宽表）
# =========================================================
raw_columns = ["变量"]
for outcome_label in OUTCOME_MAP.keys():
    for spec_name, _ in SPEC_LIST:
        raw_columns.append(f"{outcome_label}\n{spec_name}")

big_rows = []
panel_letters = ["A", "B", "C", "D", "E", "F"]

for i, moderator_label in enumerate(MODERATOR_MAP.keys()):
    panel_title = f"Panel {panel_letters[i]}  调节变量：{moderator_label}"
    row = {"变量": panel_title}
    for c in raw_columns[1:]:
        row[c] = ""
    big_rows.append(row)

    row_structure = [
        ("团队HSS浓度", "coef"),
        (moderator_label, "coef"),
        (f"团队HSS浓度 × {moderator_label}", "coef"),
        ("团队规模（对数）", "coef"),
        ("先验知识存量（对数）", "coef"),
        ("常数项", "coef"),
        ("年份固定效应", "year_fe"),
        ("arXiv主分类固定效应", "topic_fe"),
        ("观测值", "nobs"),
        ("R²", "r2"),
    ]

    for row_label, row_type in row_structure:
        row = {"变量": row_label}

        for outcome_label in OUTCOME_MAP.keys():
            for spec_name, _ in SPEC_LIST:
                col_name = f"{outcome_label}\n{spec_name}"
                res = model_store.get((moderator_label, outcome_label, spec_name), None)

                if res is None:
                    row[col_name] = ""
                    continue

                if row_type == "coef":
                    term = res["terms_map"][row_label]
                    row[col_name] = fmt_coef_cell(res["model"], term)
                elif row_type == "year_fe":
                    row[col_name] = res["year_fe"]
                elif row_type == "topic_fe":
                    row[col_name] = res["topic_fe"]
                elif row_type == "nobs":
                    row[col_name] = f"{res['nobs']:,}"
                elif row_type == "r2":
                    row[col_name] = f"{res['r2']:.4f}"

        big_rows.append(row)

    # 空行
    empty_row = {"变量": ""}
    for c in raw_columns[1:]:
        empty_row[c] = ""
    big_rows.append(empty_row)

big_table_raw_df = pd.DataFrame(big_rows, columns=raw_columns)

core_long_df = pd.DataFrame(core_rows).sort_values(
    by=["调节变量", "被解释变量", "规格", "变量"]
).reset_index(drop=True)

full_params_df = pd.DataFrame(full_rows).sort_values(
    by=["调节变量", "被解释变量", "规格", "参数名"]
).reset_index(drop=True)

model_summary_df = pd.DataFrame(summary_rows).sort_values(
    by=["调节变量", "被解释变量", "规格"]
).reset_index(drop=True)

readme_rows = [
    ["输入文件", str(CSV_PATH)],
    ["输出文件", str(OUTPUT_XLSX)],
    ["团队HSS浓度变量", BASE_VAR],
    ["控制变量1", CONTROL_VARS[0]],
    ["控制变量2", CONTROL_VARS[1]],
    ["年份固定效应变量", YEAR_FE_VAR],
    ["arXiv主分类固定效应变量", TOPIC_FE_VAR],
    ["稳健标准误", ROBUST_SE],
    ["说明", "按最终确认口径：5个知识扩散指标为因变量，4个参考文献侧指标为调节变量"],
    ["说明", "无主题FE/有主题FE使用同一批样本，便于比较系数变化"],
    ["", ""],
    ["因变量显示名", "CSV 实际列名"],
]
for k, v in OUTCOME_MAP.items():
    readme_rows.append([k, v])

readme_rows.append(["", ""])
readme_rows.append(["调节变量显示名", "CSV 实际列名"])
for k, v in MODERATOR_MAP.items():
    readme_rows.append([k, v])

readme_df = pd.DataFrame(readme_rows, columns=["项目", "内容"])


# =========================================================
# 8. 先写入 Excel
# =========================================================
with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
    readme_df.to_excel(writer, sheet_name="readme", index=False)
    big_table_raw_df.to_excel(writer, sheet_name="big_table_raw", index=False)
    core_long_df.to_excel(writer, sheet_name="core_long", index=False)
    full_params_df.to_excel(writer, sheet_name="full_params", index=False)
    model_summary_df.to_excel(writer, sheet_name="model_summary", index=False)

# =========================================================
# 9. 用 openpyxl 新建一个“漂亮版大表”
# =========================================================
wb = load_workbook(OUTPUT_XLSX)

if "big_table_pretty" in wb.sheetnames:
    del wb["big_table_pretty"]
ws = wb.create_sheet("big_table_pretty", 1)

thin = Side(border_style="thin", color="000000")
bold_font = Font(bold=True)

# 3 行表头
# A1:A3
ws["A1"] = "变量"
ws.merge_cells("A1:A3")

# 因变量顺序
outcomes = list(OUTCOME_MAP.keys())

# 行1：大类
group_spans = []
start_col = 2
current_group = None
group_start_col = start_col

for outcome_label in outcomes:
    group = OUTCOME_GROUP_MAP[outcome_label]
    if current_group is None:
        current_group = group
        group_start_col = start_col
    elif group != current_group:
        group_spans.append((current_group, group_start_col, start_col - 1))
        current_group = group
        group_start_col = start_col
    start_col += 2
group_spans.append((current_group, group_start_col, start_col - 1))

for group, c1, c2 in group_spans:
    ws.cell(1, c1).value = group
    if c1 != c2:
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)

# 行2：具体因变量
col = 2
for outcome_label in outcomes:
    ws.cell(2, col).value = outcome_label
    ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
    col += 2

# 行3：规格
col = 2
for outcome_label in outcomes:
    for spec_name, _ in SPEC_LIST:
        ws.cell(3, col).value = spec_name
        col += 1

# 数据行
data_start_row = 4

# 从 raw_df 中复制数据
for r_idx, row_data in enumerate(big_table_raw_df.itertuples(index=False), start=data_start_row):
    ws.cell(r_idx, 1).value = row_data[0]
    for c_idx in range(1, len(row_data)):
        ws.cell(r_idx, c_idx + 1).value = row_data[c_idx]

# 样式
for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

for r in range(1, 4):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).font = bold_font

for r in range(data_start_row, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if isinstance(v, str) and v.startswith("Panel "):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).font = bold_font

for cell in ws["A"]:
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 冻结
ws.freeze_panes = "B4"

# 列宽
ws.column_dimensions["A"].width = 30
for c in range(2, ws.max_column + 1):
    ws.column_dimensions[get_column_letter(c)].width = 18

# 行高
for r in range(1, 4):
    ws.row_dimensions[r].height = 28
for r in range(4, ws.max_row + 1):
    ws.row_dimensions[r].height = 34

# 其他 sheet 也稍微美化
for sheet_name in ["readme", "big_table_raw", "core_long", "full_params", "model_summary"]:
    wsx = wb[sheet_name]
    wsx.freeze_panes = "A2" if sheet_name != "big_table_raw" else "B2"
    for row in wsx.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in wsx[1]:
        cell.font = bold_font
    if sheet_name == "readme":
        wsx.column_dimensions["A"].width = 24
        wsx.column_dimensions["B"].width = 60
    else:
        for col_cells in wsx.columns:
            wsx.column_dimensions[get_column_letter(col_cells[0].column)].width = 18
    if sheet_name == "big_table_raw":
        wsx.column_dimensions["A"].width = 30
        for r in range(2, wsx.max_row + 1):
            wsx.row_dimensions[r].height = 34

wb.save(OUTPUT_XLSX)

print("=" * 70)
print("完成：已按 5 个因变量 + 4 个调节变量 + 双规格 输出结果。")
print(f"输出文件：{OUTPUT_XLSX}")
print("sheet 说明：")
print("  1) readme")
print("  2) big_table_pretty")
print("  3) big_table_raw")
print("  4) core_long")
print("  5) full_params")
print("  6) model_summary")
print("=" * 70)
